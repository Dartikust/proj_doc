"""python -m scripts.generate_dataset — 12 синтетических договоров + разметка."""

import argparse
from datetime import date
import json
from pathlib import Path

from docx import Document


def sample_record(index):
    return {
        'number': f'Д-{index:03d}/2026',
        'date': date(2026, (index - 1) % 12 + 1, min(index + 5, 28)).isoformat(),
        'customer': f'ООО «Заказчик {index}»',
        'contractor': f'ООО «Исполнитель {index}»',
        'amount': f'{10000 * index}.{25 if index % 2 else 50:02d}',
    }


def make_document(record, style):
    document = Document()
    document.add_heading(f'Договор оказания услуг № {record["number"]}', 0)
    year, month, day = record['date'].split('-')
    amount = record['amount'].replace('.', ',')
    document.add_paragraph('Учебный синтетический документ. Все организации вымышлены.')
    if style == 'prose':
        months = ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
                  'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря']
        document.add_paragraph(f'г. Новосибирск    «{day}» {months[int(month) - 1]} {year} г.')
        document.add_paragraph(f'{record["customer"]}, именуемое в дальнейшем «Заказчик», и '
                               f'{record["contractor"]}, именуемое в дальнейшем «Исполнитель», '
                               'заключили настоящий договор.')
        document.add_paragraph(f'Общая стоимость услуг составляет {amount} рублей.')
    else:
        pairs = [('Дата заключения', f'{day}.{month}.{year}'),
                 ('Заказчик', record['customer']), ('Исполнитель', record['contractor']),
                 ('Сумма договора', f'{amount} руб.')]
        if style == 'table':
            table = document.add_table(rows=0, cols=2)
            for label, value in pairs:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = label, value
        else:
            for label, value in pairs:
                document.add_paragraph(f'{label}: {value}')
    document.add_paragraph('1. Предмет договора: подготовка учебных материалов.')
    document.add_paragraph('2. Аванс составляет 1 000 рублей. Оплата до 31.12.2026.')
    document.add_paragraph('3. Стороны принимают результат по акту оказанных услуг.')
    return document


def generate(output):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    records = []
    for index in range(1, 13):
        style = ['labels', 'prose', 'table'][(index - 1) % 3]
        fields = sample_record(index)
        filename = f'contract-{index:02d}-{style}.docx'
        make_document(fields, style).save(output / filename)
        records.append({'file': filename, 'style': style, 'fields': fields})
    (output / 'expected.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    return records


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', default='data/synthetic')
    args = parser.parse_args()
    print(f'Создано {len(generate(args.output))} договоров в {args.output}')
