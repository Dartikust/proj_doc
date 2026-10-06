from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from pypdf import PdfWriter

from app import create_app
from contracts.extractor import extract_fields, normalize_text
from contracts.pipeline import process_document
from contracts.readers import DocumentError
from contracts import storage
from scripts.generate_dataset import generate, make_document, sample_record
from scripts.evaluate import evaluate


def docx_bytes(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class PipelineTests(unittest.TestCase):
    def test_all_synthetic_layouts_against_ground_truth(self):
        with tempfile.TemporaryDirectory() as directory:
            generate(directory)
            report = evaluate(directory)
        self.assertEqual(report['documents'], 12)
        self.assertEqual(report['mismatches'], [])

    def test_missing_fields_are_not_invented(self):
        document = Document()
        document.add_paragraph('Аванс 5 000 рублей. Оплата до 31.12.2026.')
        result = process_document(docx_bytes(document), '.docx')
        self.assertTrue(all(value is None for value in result.fields.values()))
        self.assertEqual(result.status, 'review')
        self.assertEqual(len(result.warnings), 5)

    def test_invalid_date_and_distracting_amount(self):
        text = ('Договор № А-7 от 31.02.2026\nЗаказчик: ООО «Ромашка»\n'
                'Исполнитель: ИП Иванов Иван Иванович\nАванс 5000 рублей.\n'
                'Стоимость договора составляет 1\u00a0234\u202f567,89 руб.\n'
                'Срок оплаты до 01.03.2026.')
        fields, warnings = extract_fields(normalize_text(text))
        self.assertEqual(fields['number'], 'А-7')
        self.assertIsNone(fields['date'])
        self.assertEqual(fields['amount'], '1234567.89')
        self.assertEqual(fields['contractor'], 'ИП Иванов Иван Иванович')
        self.assertEqual(len(warnings), 1)

    def test_empty_pdf_and_encrypted_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        blank = BytesIO()
        writer.write(blank)
        with self.assertRaisesRegex(DocumentError, 'OCR'):
            process_document(blank.getvalue(), '.pdf')
        writer.encrypt('password')
        encrypted = BytesIO()
        writer.write(encrypted)
        with self.assertRaisesRegex(DocumentError, 'паролем'):
            process_document(encrypted.getvalue(), '.pdf')

    def test_pdf_text_and_fields(self):
        fixture = Path(__file__).parent / 'fixtures' / 'contract.pdf'
        result = process_document(fixture.read_bytes(), '.pdf')
        self.assertEqual(result.fields, {
            'number': 'PDF-007', 'date': '2026-03-17', 'customer': 'ООО «Лектор»',
            'contractor': 'ИП Петров Пётр Петрович', 'amount': '125000.50',
        })
        self.assertEqual(result.warnings, [])

    def test_corrupt_and_wrong_extension(self):
        for content, extension in [(b'broken', '.pdf'), (b'broken', '.docx'), (b'text', '.txt')]:
            with self.subTest(extension=extension), self.assertRaises(DocumentError):
                process_document(content, extension)


class WebTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        root = Path(self.directory.name)
        self.config = {'TESTING': True, 'SECRET_KEY': 'test-only',
                       'DATABASE': str(root / 'registry.sqlite3'),
                       'UPLOAD_FOLDER': str(root / 'uploads')}
        self.app = create_app(self.config)
        self.client = self.app.test_client()
        self.client.get('/')
        with self.client.session_transaction() as session:
            self.token = session['csrf_token']

    def tearDown(self):
        self.directory.cleanup()

    def upload(self, content, filename):
        return self.client.post('/upload', data={'document': (BytesIO(content), filename),
                                                'csrf_token': self.token})

    def test_upload_detail_registry_download_and_restart(self):
        expected = sample_record(1)
        content = docx_bytes(make_document(expected, 'labels'))
        response = self.upload(content, '../../русский-договор.docx')
        self.assertEqual(response.status_code, 302)
        detail = self.client.get(response.location)
        self.assertEqual(detail.status_code, 200)
        self.assertIn(expected['customer'], detail.get_data(as_text=True))
        self.assertIn(expected['amount'], detail.get_data(as_text=True))
        record = storage.get_document(self.config['DATABASE'], 1)
        self.assertEqual(record['original_name'], 'русский-договор.docx')
        self.assertNotIn('..', record['stored_name'])
        download = self.client.get('/documents/1/download')
        self.assertEqual(download.data, content)
        download.close()
        second_client = create_app(self.config).test_client()
        registry = second_client.get('/').get_data(as_text=True)
        self.assertIn(expected['number'], registry)

    def test_failed_upload_does_not_create_record_or_file(self):
        for content, filename in [(b'bad', 'broken.docx'), (b'bad', 'broken.pdf'), (b'bad', 'other.exe')]:
            with self.subTest(filename=filename):
                response = self.upload(content, filename)
                self.assertEqual(response.status_code, 302)
        self.assertEqual(storage.list_documents(self.config['DATABASE']), [])
        self.assertEqual(list(Path(self.config['UPLOAD_FOLDER']).iterdir()), [])

    def test_missing_fields_are_visible(self):
        document = Document()
        document.add_paragraph('Договор № 15')
        response = self.upload(docx_bytes(document), 'partial.docx')
        detail = self.client.get(response.location).get_data(as_text=True)
        self.assertIn('Не все поля удалось извлечь', detail)
        record = storage.get_document(self.config['DATABASE'], 1)
        self.assertEqual(record['status'], 'review')
        self.assertEqual(len(json.loads(record['warnings'])), 4)

    def test_sample_is_valid_and_missing_document_is_404(self):
        response = self.client.get('/sample')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(process_document(response.data, '.docx').fields, sample_record(1))
        response.close()
        self.assertEqual(self.client.get('/documents/999').status_code, 404)
        self.assertEqual(self.client.get('/documents/999/download').status_code, 404)

    def test_csrf_and_size_limit(self):
        self.assertEqual(self.client.post('/upload').status_code, 400)
        self.app.config['MAX_CONTENT_LENGTH'] = 100
        response = self.upload(b'x' * 1000, 'large.pdf')
        self.assertEqual(response.status_code, 413)
        self.assertIn('слишком большой', response.get_data(as_text=True))


if __name__ == '__main__':
    unittest.main()
